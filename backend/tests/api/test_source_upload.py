"""
R5.1: uploading a document to the approved-source library. The text is
extracted for review; the source is only created by the usual
create-as-draft call, then approved as before.
"""

import io

from docx import Document

from tests.conftest import ok

POLICY = (
    "Customer Due Diligence Policy\n\n"
    "Beneficial owners holding 25% or more must be identified and verified before onboarding.\n\n"
    "Enhanced due diligence applies to customers in high-risk third countries."
)


def _upload(client, headers, name, data, mime="text/plain"):
    return client.post("/api/sources/extract", files={"file": (name, data, mime)}, headers=headers)


def test_library_admin_uploads_a_text_document_and_saves_it_as_a_draft(client, auth):
    before = len(ok(client.get("/api/sources", headers=auth("admin"))))
    extracted = ok(_upload(client, auth("admin"), "cdd_policy-v3.txt", POLICY.encode()))
    assert extracted["filename"] == "cdd_policy-v3.txt"
    assert extracted["title"] == "cdd policy v3"
    assert "Beneficial owners holding 25%" in extracted["content"]
    assert extracted["characters"] == len(extracted["content"])
    # Extraction alone stores nothing.
    assert len(ok(client.get("/api/sources", headers=auth("admin")))) == before

    source = ok(
        client.post(
            "/api/sources",
            json={
                "title": "CDD Policy",
                "source_type": "INTERNAL_POLICY",
                "version": "3",
                "reference": extracted["filename"],
                "content": extracted["content"],
            },
            headers=auth("admin"),
        ),
        201,
    )
    assert source["status"] == "DRAFT"
    ok(client.post(f"/api/sources/{source['id']}/approve", json={"reason": "Reviewed."}, headers=auth("admin")))
    hits = ok(client.get("/api/sources/search?q=beneficial owners", headers=auth("analyst")))
    assert any(hit["source_id"] == source["id"] for hit in hits)


def test_word_documents_are_read(client, auth):
    document = Document()
    document.add_paragraph("Vendor Oversight Standard")
    document.add_paragraph("Processors must provide a SOC 2 Type II report every year.")
    buffer = io.BytesIO()
    document.save(buffer)
    extracted = ok(
        _upload(
            client, auth("admin"), "vendor-standard.docx", buffer.getvalue(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    )
    assert "SOC 2 Type II" in extracted["content"]


def test_only_library_admins_may_upload(client, auth):
    for who in ("analyst", "manager", "owner"):
        assert _upload(client, auth(who), "policy.txt", POLICY.encode()).status_code == 403


def _create_with_file(client, headers, name, data, **fields):
    form = {"title": "Sanctions Policy", "source_type": "INTERNAL_POLICY", "version": "1", "content": POLICY, **fields}
    return client.post("/api/sources/with-file", data=form, files={"file": (name, data, "text/plain")}, headers=headers)


def test_the_original_document_is_kept_encrypted_and_downloadable(client, auth):
    import hashlib

    from app.database import SessionLocal
    from app.file_processing.storage import encryption_enabled, is_encrypted
    from app.models.approved_source import ApprovedSource
    from app.models.audit_event import AuditEvent

    original = POLICY.encode() + b"\n\nOriginal-only footer line."
    source = ok(_create_with_file(client, auth("admin"), "cdd-policy.txt", original, reference="POL-CDD"), 201)
    assert source["status"] == "DRAFT" and source["has_file"] is True
    assert source["original_filename"] == "cdd-policy.txt" and source["file_size"] == len(original)
    assert source["file_sha256"] == hashlib.sha256(original).hexdigest()
    assert source["file_content_type"] == "text/plain"
    assert "file_path" not in source  # the storage path is never exposed
    assert source["content"] == POLICY  # the reviewed text, not the raw file

    db = SessionLocal()
    try:
        stored = open(db.get(ApprovedSource, source["id"]).file_path, "rb").read()
    finally:
        db.close()
    if encryption_enabled():
        assert is_encrypted(stored) and original not in stored

    url = f"/api/sources/{source['id']}/file"
    # A draft's file is hidden from non-admins, like the draft itself.
    assert client.get(url, headers=auth("analyst")).status_code == 404
    admin_copy = client.get(url, headers=auth("admin"))
    assert admin_copy.status_code == 200 and admin_copy.content == original
    assert "attachment" in admin_copy.headers["content-disposition"] and "cdd-policy.txt" in admin_copy.headers["content-disposition"]

    ok(client.post(f"/api/sources/{source['id']}/approve", json={"reason": "Reviewed."}, headers=auth("admin")))
    analyst_copy = client.get(url, headers=auth("analyst"))
    assert analyst_copy.status_code == 200 and analyst_copy.content == original

    db = SessionLocal()
    try:
        downloads = db.query(AuditEvent).filter(
            AuditEvent.action == "DOCUMENT_DOWNLOADED", AuditEvent.details.like(f"%source #{source['id']} %")
        ).count()
    finally:
        db.close()
    assert downloads == 2

    # The file can't be swapped: draft edits don't touch it, and an approved source isn't editable.
    assert client.patch(f"/api/sources/{source['id']}", json={"title": "x"}, headers=auth("admin")).status_code == 409


def test_pasted_sources_have_no_file_and_non_admins_cannot_upload(client, auth):
    pasted = ok(
        client.post(
            "/api/sources",
            json={"title": "Pasted", "source_type": "PROCEDURE", "version": "1", "content": POLICY},
            headers=auth("admin"),
        ),
        201,
    )
    assert pasted["has_file"] is False and pasted["original_filename"] is None
    assert client.get(f"/api/sources/{pasted['id']}/file", headers=auth("admin")).status_code == 404
    assert _create_with_file(client, auth("analyst"), "x.txt", b"text").status_code == 403
    assert _create_with_file(client, auth("admin"), "x.exe", b"MZ").status_code == 400
    assert _create_with_file(client, auth("admin"), "x.txt", b"text", version="  ").status_code == 422


def test_unsupported_and_empty_files_are_refused(client, auth):
    assert _upload(client, auth("admin"), "macro.exe", b"MZ...", "application/octet-stream").status_code == 400
    assert _upload(client, auth("admin"), "empty.txt", b"").status_code == 400
    blank = _upload(client, auth("admin"), "blank.txt", b"   \n  ")
    assert blank.status_code == 422 and "No readable text" in blank.json()["detail"]
