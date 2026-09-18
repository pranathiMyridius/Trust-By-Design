from pathlib import Path
import csv
import io

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {
    ".docx",
    ".xlsx",
    ".pdf",
    ".txt",
    ".csv",
}


def extract_text(filename: str, file_content: bytes) -> str:
    """
    Extract readable text from supported document types.

    Supported:
    - DOCX
    - XLSX
    - PDF
    - TXT
    - CSV
    """

    extension = Path(filename).suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported file type: {extension}"
        )

    if extension == ".docx":
        return _extract_docx(file_content)

    if extension == ".xlsx":
        return _extract_xlsx(file_content)

    if extension == ".pdf":
        return _extract_pdf(file_content)

    if extension == ".txt":
        return _extract_txt(file_content)

    if extension == ".csv":
        return _extract_csv(file_content)

    raise ValueError(
        f"Unsupported file type: {extension}"
    )


def _extract_docx(file_content: bytes) -> str:
    """Extract paragraphs and table content from Word."""

    document = Document(
        io.BytesIO(file_content)
    )

    sections = []

    # Extract paragraphs
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            sections.append(text)

    # Extract tables
    for table in document.tables:
        for row in table.rows:
            row_values = []

            for cell in row.cells:
                value = cell.text.strip()

                if value:
                    row_values.append(value)

            if row_values:
                sections.append(" | ".join(row_values))

    return "\n".join(sections)


def _extract_xlsx(file_content: bytes) -> str:
    """Extract values from all Excel worksheets."""

    workbook = load_workbook(
        filename=io.BytesIO(file_content),
        read_only=True,
        data_only=True,
    )

    sections = []

    for worksheet in workbook.worksheets:

        sections.append(
            f"[Sheet: {worksheet.title}]"
        )

        for row in worksheet.iter_rows(
            values_only=True
        ):
            values = []

            for value in row:
                if value is not None:
                    values.append(str(value).strip())

            if values:
                sections.append(
                    " | ".join(values)
                )

    return "\n".join(sections)


def _extract_pdf(file_content: bytes) -> str:
    """Extract text from all PDF pages."""

    reader = PdfReader(
        io.BytesIO(file_content)
    )

    sections = []

    for page_number, page in enumerate(
        reader.pages,
        start=1,
    ):
        text = page.extract_text()

        if text:
            sections.append(
                f"[Page {page_number}]"
            )
            sections.append(
                text.strip()
            )

    return "\n".join(sections)


def _extract_txt(file_content: bytes) -> str:
    """Extract text from a plain text file."""

    return file_content.decode(
        "utf-8",
        errors="replace",
    ).strip()


def _extract_csv(file_content: bytes) -> str:
    """Extract CSV rows as readable text."""

    decoded = file_content.decode(
        "utf-8",
        errors="replace",
    )

    reader = csv.reader(
        io.StringIO(decoded)
    )

    sections = []

    for row in reader:
        values = [
            value.strip()
            for value in row
            if value.strip()
        ]

        if values:
            sections.append(
                " | ".join(values)
            )

    return "\n".join(sections)